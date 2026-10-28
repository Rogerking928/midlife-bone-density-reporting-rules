"""圖表本（Word）＋ 700 dpi PNG ＋ Source_data.xlsx，全部放同一個交付資料夾。

版式照 Bcc 那篇（使用者 2026-10-23 指定「參考 BCG 那篇」＝Bcc；docx 部分沿用墨西哥剖腹產那篇的 book.py）：
Letter 橫向、左右 1.0／上下 0.9 吋、每項新起一頁；標題 TNR 11.5 pt 粗體；
表格滿版 8 pt、表頭粗體、三條黑線；表註 9 pt、圖說 9.5 pt，正體。
順序：主文表 → 主文圖 → eTable → eFigure，一份 Word。**只交 PNG，不產 PDF。**

用法：python3 code/book.py [輸出資料夾]
"""
import os
import json, os, re, shutil, sys
import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from openpyxl import Workbook
from openpyxl.styles import Font as XFont, Alignment
from pathlib import Path
OUT = Path(__file__).resolve().parent.parent / "out"

FONT = "Times New Roman"
PAGE_W, PAGE_H = Inches(11.0), Inches(8.5)
MARGIN_LR, MARGIN_TB = Inches(1.0), Inches(0.9)
BOX_W, BOX_H = 8.2, 5.7
MAX_SCALE = 1.5
DEFAULT_DEST = os.environ.get("DELIVERABLE_DIR", (os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "deliverable")))


# ------------------------------------------------------------------ docx
def run(p, text, size=None, bold=None, italic=None):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size: r.font.size = Pt(size)
    if bold is not None: r.bold = bold
    if italic is not None: r.italic = italic
    return r


def run_sci(p, text, size, bold=None):
    """<i>…</i> 包住的字設斜體，其餘正體。"""
    for part in re.split(r"(<i>.*?</i>|<sub>.*?</sub>)", text):
        if part.startswith("<i>"):
            run(p, part[3:-4], size, bold=bold, italic=True)
        elif part.startswith("<sub>"):
            run(p, part[5:-6], size, bold=bold, italic=False).font.subscript = True
        elif part:
            run(p, part, size, bold=bold, italic=False)


def para(doc, before=0, after=0, align=None, page_break=False, keep=False):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), 1.0
    f.page_break_before, f.keep_with_next, f.keep_together = page_break, keep, keep
    if align is not None: p.alignment = align
    return p


def borders(cell, *edges):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")): tcPr.remove(old)
    b = OxmlElement("w:tcBorders")
    for e in edges:
        el = OxmlElement(f"w:{e}")
        for k, v in (("val", "single"), ("sz", "8"), ("space", "0"), ("color", "000000")):
            el.set(qn(f"w:{k}"), v)
        b.append(el)
    tcPr.append(b)


def add_table(doc, item, first):
    p = para(doc, after=4, page_break=not first, keep=True)
    label, rest = item["caption"].split(". ", 1)
    run(p, label + ". ", 11.5, bold=True); run(p, rest, 11.5, bold=True)
    rows = item["rows"]
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    first_w = item.get("first_w") or (min(3.6, BOX_W * 0.34) if ncol > 2 else BOX_W * 0.55)
    widths = item.get("widths") or [first_w] + [(BOX_W - first_w) / (ncol - 1)] * (ncol - 1)
    for i, row in enumerate(rows):
        is_group = i > 0 and all(c == "" for c in row[1:])
        for j, text in enumerate(row):
            c = t.cell(i, j)
            c.width = Inches(widths[j])
            cp = c.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0); cp.paragraph_format.line_spacing = 1.0
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run_sci(cp, str(text), 8, bold=(i == 0 or (is_group and j == 0)))
    last = len(rows) - 1
    for j in range(ncol):
        borders(t.cell(0, j), "top", "bottom")
        if last >= 1: borders(t.cell(1, j), *(("top", "bottom") if last == 1 else ("top",)))
        if last > 1: borders(t.cell(last, j), "bottom")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    if item.get("note"):
        n = para(doc, before=3)
        run_sci(n, item["note"], 9)


def add_figure(doc, item, first):
    p = para(doc, after=4, page_break=not first, keep=True)
    run(p, item["caption"], 11.5, bold=True)
    if item.get("note"):
        n = para(doc, after=6, keep=True)
        run_sci(n, item["note"], 9.5)
    w_px, h_px = Image.open(item["path"]).size
    w_in, h_in = w_px / 700, h_px / 700
    # 可用高度＝版心 6.7 吋 − 標題 − 圖說（9.5 pt 約每行 150 字、0.17 吋）− 安全邊
    lines = -(-len(item.get("note", "")) // 150)
    avail = min(BOX_H, (PAGE_H - 2 * MARGIN_TB) / 914400 - 0.3 - 0.17 * lines - 0.15)
    s = min(MAX_SCALE, BOX_W / w_in, avail / h_in)
    ip = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    ip.add_run().add_picture(str(item["path"]), width=Inches(w_in * s), height=Inches(h_in * s))


def add_text(doc, item, first):
    """Supporting Methods：標題＋小節標題＋段落（10 pt）。"""
    p = para(doc, after=6, page_break=not first, keep=True)
    run(p, item["caption"], 11.5, bold=True)
    for head, body in item["sections"]:
        h = para(doc, before=6, after=2, keep=True)
        run(h, head, 10, bold=True)
        for par in body:
            q = para(doc, after=4)
            run_sci(q, par, 10)


def build(items, path, supp=()):
    """一份 Word：主文的表、主文的圖，接著補充的表、補充的圖（使用者 2026-10-04：圖表 Word 一個就好）。"""
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(8)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = PAGE_W, PAGE_H
    sec.left_margin = sec.right_margin = MARGIN_LR
    sec.top_margin = sec.bottom_margin = MARGIN_TB
    tabs = [i for i in items if i["kind"] == "table"]
    figs = [i for i in items if i["kind"] == "figure"]
    stext = [i for i in supp if i["kind"] == "text"]
    stabs = [i for i in supp if i["kind"] == "table"]
    sfigs = [i for i in supp if i["kind"] == "figure"]
    for k, it in enumerate(tabs + figs + stext + stabs + sfigs):
        {"table": add_table, "figure": add_figure, "text": add_text}[it["kind"]](doc, it, k == 0)
    doc.save(path)
    return len(tabs) + len(stabs), len(figs) + len(sfigs)




# ------------------------------------------------------------------ 共用
NUMCELL = re.compile(r"^[\s\d,.+\-−()/%]*(?:to[\s\d,.+\-−()/%]*)*$")


def split_cell(text):
    """'1.38 (1.34 to 1.42)' → [1.38, 1.34, 1.42]；'425 / 37' → [425, 37]；文字格回傳 None。"""
    t = str(text).strip()
    if not t or not re.search(r"\d", t) or re.fullmatch(r"\d{4}-\d{2}-\d{2}", t) or not NUMCELL.match(t):
        return None
    return [Dec(x.replace(",", "").replace("−", "-")) for x in re.findall(r"[+\-−]?\d[\d,]*(?:\.\d+)?", t)]


class Dec(float):
    """記得原字串小數位數的 float；寫 Excel 時用來設 number_format。"""
    def __new__(cls, txt):
        o = float.__new__(cls, txt); o.nd = len(txt.split(".")[1]) if "." in txt else 0
        return o


def numeric_table(rows):
    """表格工作表存數值：每欄依最多的數字個數拆成子欄（點估計／下限／上限，或以 / 分隔的各值）。"""
    head, body = rows[0], rows[1:]
    out_cols, data = [], [[] for _ in body]
    for j, h in enumerate(head):
        parsed = [split_cell(r[j]) if j > 0 else None for r in body]
        k = max((len(p) for p in parsed if p), default=0)
        if j == 0 or k == 0:
            out_cols.append(h)
            for i, r in enumerate(body):
                data[i].append(str(r[j]))
            continue
        cells = [str(r[j]) for r in body]
        if any(" to " in c for c in cells) and k == 3:
            subs = [h, h + " - lower 95% limit", h + " - upper 95% limit"]
        elif ":" in h and "/" in h and len(h.split(":")[-1].split("/")) == k:
            pre = h.split(":")[0] + ": " if ":" in h else ""
            subs = [pre + x.strip() for x in h.split(":")[-1].split("/")]
        else:
            subs = [h] + [f"{h} - value {q + 1}" for q in range(1, k)]
        out_cols += subs
        for i, p in enumerate(parsed):
            if p is None:
                data[i] += [cells[i] if cells[i] else None] + [None] * (k - 1)
            else:
                data[i] += p + [None] * (k - len(p))
    return pd.DataFrame(data, columns=out_cols, dtype=object)





